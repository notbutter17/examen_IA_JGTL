"""Genera el artículo editable de seis páginas a partir de resultados calculados."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];O=R/'resultados';F=O/'figuras'
def load(n):return json.loads((O/(n+'.json')).read_text(encoding='utf8'))
a=load('auditoria_datos');m=load('resumen_modelos');c=load('resumen_clustering');v=load('resumen_desventaja');sh=load('resumen_shap')
cv=pd.read_csv(O/'comparacion_validacion.csv');tests=pd.read_csv(O/'comparacion_test.csv').set_index('modelo');groups=pd.read_csv(O/'doble_desventaja.csv');prof=pd.read_csv(O/'perfiles_clusters.csv')
doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.03
for name in ['Heading 1','Heading 2']:
 doc.styles[name].font.size=Pt(12 if name=='Heading 1' else 11);doc.styles[name].paragraph_format.space_before=Pt(7);doc.styles[name].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.bold=True;doc.styles['Caption'].font.size=Pt(8.5)
doc.core_properties.author='Jose Guillermo Ticona Lipa';doc.core_properties.title='Doble desventaja digital y contexto docente asociados al rezago escolar en el Perú'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Jose Guillermo Ticona Lipa  |  ');run.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
source=[]
next_page=False
doc.core_properties.comments=''
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
def p(t,bold=False):
 source.append(t);x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;r=x.add_run(t);r.bold=bold;return x

def h(t):
 global next_page
 source.append('\n'+t);pp=doc.add_heading(t,1)
 if next_page:pp.paragraph_format.page_break_before=True;next_page=False
def img(n,caption,width=17):
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(1);x.add_run().add_picture(str(F/(n+'.png')),width=Cm(width));x=doc.add_paragraph(caption,'Caption');x.paragraph_format.space_after=Pt(5)
def pair(n1,n2,caption,width=8.3):
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(1)
 for n in [n1,n2]:x.add_run().add_picture(str(F/(n+'.png')),width=Cm(width))
 doc.add_paragraph(caption,'Caption')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for cell,value in zip(t.rows[0].cells,headers):cell.text=str(value)
 for row in rows:
  for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
 for i,row in enumerate(t.rows):
  for j,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[j])
   cell.vertical_alignment=1
   tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders);shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E5EDF1' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF'));tcpr.append(shd)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.space_before=Pt(3)
    for rr in pp.runs:rr.font.size=Pt(8.5);rr.bold=i==0
  if i==0:
   repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
 return t

def page():
 global next_page
 next_page=True
# PAGE 1
x=doc.add_paragraph('Doble desventaja digital y contexto docente asociados al rezago escolar en el Perú','Title');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
x=doc.add_paragraph('Jose Guillermo Ticona Lipa\nProyecto de Inteligencia Artificial  |  ENAHO y ESCALE 2025');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
h('Resumen')
p(f'Se estudia si la ausencia de conexión doméstica y un ratio contextual alto de alumnos por docente coinciden con mayor desfase de edad y grado. Se integraron {a["filas"]:,} estudiantes de 6 a 17 años del corte 2025 del panel ENAHO con servicios educativos activos de ESCALE por distrito y área. Se compararon K-Means y mezclas gaussianas, y regresión logística frente a tres ensambles con validación por conglomerado. El XGBoost seleccionado obtuvo ROC-AUC {m["test"]["ROC_AUC"]:.3f}, AP {m["test"]["PR_AUC_AP"]:.3f} y F1 macro {m["test"]["F1_macro"]:.3f} en test. La celda de doble desventaja contiene 24 casos y una tasa de desfase de 58,3 %, con gran incertidumbre. El contraste aditivo incluye cero en su intervalo bootstrap. Los resultados no confirman una acumulación generalizable y no justifican decisiones individuales. SHAP describe asociaciones del modelo, sin interpretación causal.')
p('Palabras clave  brecha digital, contexto docente, desfase edad grado, validación agrupada, perfiles de recursos, SHAP.')
h('1 Introducción y pregunta de investigación')
p('Las condiciones tecnológicas del hogar y los recursos educativos del territorio pueden coincidir sin ser equivalentes. Un distrito con más alumnos por docente no identifica la escuela concreta del estudiante ni mide su tamaño de aula. Por ello, el análisis distingue exposición contextual, clasificación individual y descripción de perfiles.')
p('La pregunta es si la coexistencia de ausencia de conexión doméstica y ratio docente alto presenta una asociación adicional con el desfase escolar. Se examinan también estudiantes sin desfase dentro de esa combinación. Estos casos se denominan trayectorias acordes a la edad en la medición disponible; no prueban resiliencia psicológica ni identifican sus causas.')
p('El objetivo general es construir un análisis reproducible que contraste esa pregunta mediante integración relacional, clustering de recursos y clasificación explicable. Se anticipa la hipótesis de mayor desfase en la doble desventaja, pero se conserva la posibilidad de resultados negativos o inconclusos.')
img('diagrama_cruces','Figura 1. Integración por persona, hogar y distrito con área. Las flechas indican dependencias de integración.',16)
# PAGE 2
page();h('2 Datos y decisiones metodológicas')
p('Se procesaron las selecciones originales 2025 de los módulos 100, 200, 300 y Sumaria incluidos en los archivos panel 2021–2025 [1]. Se verificó UTF-8 con BOM en estos CSV; no se impuso latin1 solo por la advertencia general de la guía. Las llaves se leyeron como texto, se conservaron sus ceros y se normalizaron espacios. Los hashes SHA-256 registran la procedencia y permiten reconstruir las selecciones desde los ZIP originales.')
flow=pd.read_csv(O/'flujo_poblacion.csv');table(['Filtro acumulativo','Registros'],[[r.etapa,f'{r.n:,}'] for _,r in flow.iterrows()],[13,3.5])
p('Son residentes habituales los miembros no ausentes durante 30 días o más, y los no miembros presentes durante al menos 30 días, según P204–P206. Se retienen matriculados en primaria o secundaria con grado válido. La unión 200–300 es uno a uno por CONGLOME, VIVIENDA, HOGAR y CODPERSO; vivienda y Sumaria se vinculan muchos a uno por hogar. No se detectaron duplicados en esas llaves ni pérdidas en los tres cruces para la cohorte final.')
p(f'ESCALE aporta {a["servicios_primaria_secundaria_activos"]:,} servicios activos B0/F0 [2]. El ratio contextual es la suma de TALUM dividida por la suma de TDOC, sobre servicios con alumnos válidos y docentes positivos, por CODGEO y área. Se calcula también proporción pública y coordenadas medias válidas. Padlocal verifica CODLOCAL; aparece una discrepancia geográfica, que se registra conservando CODGEO del padrón. La cobertura contextual es {a["cobertura_contexto"]:.1%}; los 369 estudiantes no vinculados permanecen con faltantes.')
h('2 1 Definición del resultado y las exposiciones')
p('En primaria el grado actual procede de P308C; en secundaria se suma seis al año P308B. La edad normativa aproximada es grado acumulado + 5. El objetivo principal vale uno cuando la edad declarada supera esa edad, conforme a la opción amplia de la guía. No hay fecha de nacimiento usable para reconstruir la edad al 31 de marzo [3]. Se reporta por separado sensibilidad de dos o más años, cercana al umbral de atraso de ESCALE [4]; ambas definiciones no son intercambiables.')
p('P1144 y P1121 son selecciones múltiples: 1 identifica la opción seleccionada y 0 la opción no seleccionada. P314A usa 1 sí y 2 no. Se preservan faltantes. P314D resume frecuencia ordinal, no horas ni competencias. Ingreso y gasto anuales se dividen entre miembros del hogar y 12; se aplica log1p. Los códigos de lengua y pobreza se mantienen como categorías.')
h('2 2 Separación y prevención de fuga')
p(f'GroupShuffleSplit reserva {m["n_test"]:,} casos de {m["conglomerados_test"]} conglomerados para test y {m["n_train"]:,} casos para entrenamiento. Cinco folds GroupKFold, todos disjuntos por CONGLOME, ajustan medianas, escalado, categorías y pesos de balance únicamente dentro de cada train [5]. Edad, grado, normativa, desfase, identificadores, factores y clusters se excluyen de X. Se usa semilla 42 y umbral 0,5.')
p('FACTOR07 se conserva como peso del hogar; no se utiliza como si fuera individual. FACPOB07 se emplea en sensibilidad de evaluación. Ninguno convierte automáticamente el corte del panel en una estimación nacional. Padron y Padlocal suministrados no contienen altitud ni medidas completas de conexión escolar: se declaran ausencias y no se inventan indicadores.')
# PAGE 3
page();h('3 Acumulación de desventajas y perfiles')
p(f'Ratio alto significa ratio contextual mayor o igual a {v["umbral_ratio_train_P75"]:.2f}, percentil 75 calculado en train. No es un estándar de sobrecarga MINEDU. La clasificación 2 × 2 utiliza conexión doméstica y este umbral; el test no fija cortes ni modifica la hipótesis. Las tasas son descriptivas de la muestra y excluyen filas sin ambas exposiciones.')
img('doble_desventaja','Figura 2. Tasas en test e intervalos del 95 % por bootstrap de conglomerados dentro de cada celda.',15)
p(f'La combinación de ambas desventajas presenta 14 de 24 estudiantes con desfase y 10 sin él. La diferencia de diferencias de las cuatro tasas es {100*v["contraste_aditivo_test"]:.1f} puntos porcentuales, IC95 % [{100*v["IC95_contraste_bootstrap_conglomerados"][0]:.1f}; {100*v["IC95_contraste_bootstrap_conglomerados"][1]:.1f}], con 1 000 réplicas agrupadas. Este intervalo incluye cero. El resultado es compatible con diversas magnitudes y no demuestra sinergia ni protección del acceso a Internet. La celda pequeña obliga a evitar recomendaciones dirigidas a individuos.')
h('3 1 Clustering de recursos')
p(f'Se compararon K-Means y GMM para k de 2 a 6, con imputación y escalado ajustados en train. La selección por silueta sobre una muestra fija de 1 800 registros favoreció {c["ganador"]["algoritmo"]} con k = {int(c["ganador"]["k"])}, silueta {c["ganador"]["silueta"]:.3f}, Davies–Bouldin {c["ganador"]["Davies_Bouldin"]:.3f} y Calinski–Harabasz {c["ganador"]["Calinski_Harabasz"]:.1f}. El máximo se alcanza en el límite explorado: seis no se presenta como número óptimo universal.')
rows=[]
for _,r in prof.iterrows():rows.append([int(r.cluster),int(r.n),f'{r.rural:.0%}',f'{r.internet_hogar:.0%}',f'{r.ratio_alumno_docente:.1f}',f'{r.rezago:.1%}'])
table(['Perfil','n','Rural','Internet','Ratio','Desfase'],rows,[1.5,2.2,2.3,2.3,2.3,2.6])
p('Los nombres completos describen ruralidad, conexión y ratio, evitando categorías de valor sobre los estudiantes. Se proponen verificaciones territoriales de acceso para perfiles con poca conexión y de acompañamiento para perfiles con conexión pero poco uso. Son hipótesis de intervención, no conclusiones de impacto.')
img('pca_siluetas','Figura 3. PCA y siluetas de train. Las dos componentes resumen el 53,3 % de la varianza; no representan toda la estructura.',16)
# PAGE 4
page();h('4 Comparación predictiva y análisis de errores')
p('La logística es la línea base y se compara con Random Forest, XGBoost y LightGBM. Se prueban dos configuraciones por familia con los mismos cinco folds; 40 ajustes en total. C se prueba en 0,1 y 1; mínimo de hoja en 5 y 20; profundidad en 2 y 4; hojas LightGBM en 7 y 15. La mayor AP media de validación elige familia y parámetros. El test se consulta después de cerrar esa selección.')
rows=[]
for _,r in cv.iterrows():t=tests.loc[r.modelo];rows.append([r.modelo,f'{r.CV_PR_AUC_AP:.3f}',f'{t.ROC_AUC:.3f}',f'{t.PR_AUC_AP:.3f}',f'{t.F1_macro:.3f}'])
table(['Modelo','AP CV','AUC test','AP test','F1 macro'],rows,[4.3,2.5,3,3,3])
p(f'XGBoost alcanza AUC {m["test"]["ROC_AUC"]:.3f}, IC95 % [{m["IC95_bootstrap_conglomerados"]["ROC_AUC"][0]:.3f}; {m["IC95_bootstrap_conglomerados"]["ROC_AUC"][1]:.3f}]. AP se calcula como average precision, no como área trapezoidal. La referencia de AP es la prevalencia test {m["prevalencia_test"]:.3f}; la mejora es pequeña. Se entregan todas las corridas para evitar presentar solo el mejor resultado.')
img('roc_pr','Figura 4. Curvas ROC y Precision Recall de modelos seleccionados por validación. Línea horizontal PR igual a prevalencia.',16.5)
img('confusion_calibracion','Figura 5. Matriz de confusión y calibración del ganador con umbral fijo 0,5.',16.5)
p('Se observan 771 falsos negativos y 987 falsos positivos. Los primeros omiten estudiantes con desfase; los segundos generan alertas innecesarias. El balance de clases no asegura calibración: las probabilidades se concentran cerca de 0,5. El Brier de 0,248 supera el de predecir siempre la prevalencia test, aproximadamente 0,237; este último es una referencia descriptiva calculada a posteriori, no un modelo seleccionado.')
p(f'La separación adicional por distrito obtiene AUC {m["sensibilidad_distrito"]["ROC_AUC"]:.3f}. Con FACPOB07, AUC {m["sensibilidad_FACPOB07"]["ROC_AUC"]:.3f}. Al redefinir atraso como dos o más años, el mismo tipo de modelo con parámetros fijados obtiene AUC {m["sensibilidad_atraso_2_anios"]["ROC_AUC"]:.3f} y AP {m["sensibilidad_atraso_2_anios"]["PR_AUC_AP"]:.3f}; esa tarea distinta no reemplaza el resultado principal ni rescata su rendimiento.')
# PAGE 5
page();h('5 Explicabilidad global y dependencias')
p('Se aplica SHAP al XGBoost ganador [6], usando 80 observaciones de entrenamiento como fondo y 350 de test para explicación. Las contribuciones se expresan en margen del modelo, no en puntos porcentuales de probabilidad. El error máximo de reconstrucción de la salida es menor que 0,000001. Las variables mostradas se restituyen a sus unidades originales después del escalado; los logaritmos conservan su escala documentada.')
img('shap_global','Figura 6. Diez variables con mayor contribución absoluta media en la muestra explicada.',14.6)
p('Latitud, proporción pública, frecuencia TIC e ingreso encabezan las contribuciones. El patrón refleja asociaciones dentro de este modelo débil. Que una coordenada sea importante sugiere estructura territorial compartida o variables omitidas; no convierte la localización en causa del desfase. La magnitud SHAP tampoco indica que una intervención sobre esa variable produciría el cambio predicho.')
pair('shap_dependencia_internet','shap_dependencia_ratio','Figura 7. Dependencias de conexión doméstica y ratio alumno/docente, coloreadas por quechua y ruralidad respectivamente.')
p('La dependencia de Internet toma dos valores y no representa una curva dosis respuesta. Los colores permiten examinar heterogeneidad por contexto, pero no constituyen valores formales de interacción SHAP. En el ratio, la dispersión vertical indica que otros predictores también intervienen. La escasez de casos sin conexión reduce la estabilidad de los patrones; un contraste visual no reemplaza una prueba de interacción.')
p('La explicación lingüística se interpreta con prudencia: la codificación one hot conserva categorías, pero el tamaño de cada grupo varía. La ausencia de una lengua en las diez variables principales no demuestra ausencia de desigualdad. Tampoco corresponde ordenar grupos étnicos por su contribución al riesgo.')
# PAGE 6
page();h('5 1 Dos predicciones individuales')
pair('shap_caso_alto','shap_caso_bajo','Figura 8. Waterfalls de dos casos extremos correctamente clasificados de test.',8.1)
p(f'El caso alto tiene clase real 1 y probabilidad {sh["casos"]["alto"]["probabilidad"]:.3f}; el bajo tiene clase real 0 y probabilidad {sh["casos"]["bajo"]["probabilidad"]:.3f}. Son extremos relativos del conjunto explicado, no probabilidades cercanas a certeza ni casos representativos. Las barras descomponen el margen alrededor del valor base. Estos ejemplos ilustran el funcionamiento del modelo y no autorizan diagnósticos escolares.')
h('6 Discusión y conclusiones')
p('La hipótesis de acumulación recibe una señal descriptiva positiva en una celda pequeña, pero su contraste agrupado es impreciso e incluye cero. No se confirma una asociación adicional estable. Los diez casos sin desfase bajo doble desventaja abren preguntas sobre acompañamiento, acceso alternativo y trayectoria previa que estos datos no permiten responder; no se atribuyen mecanismos que no fueron medidos.')
p('Los perfiles principales se construyen sin desfase para evitar que la comparación de rezago sea definitoria. Como sensibilidad exigida por las variables sugeridas en la guía, se repite la comparación de ambos algoritmos incluyendo años de desfase. Se entrega la tabla alternativa y su concordancia con los perfiles de recursos; ningún cluster entra en el clasificador. Toda interpretación del perfil alternativo reconoce la participación del resultado en su construcción.')
p('Las limitaciones incluyen corte panel sin equivalencia certificada con la anual, edad declarada sin referencia al 31 de marzo, ratio territorial que no mide carga individual, ausencia de altitud y de detalles completos de equipamiento escolar, y pocos casos en la doble desventaja. La separación por conglomerado reduce fuga por vecindario, pero no garantiza independencia espacial total. SHAP explica el ajuste y no identifica efectos causales.')
p('Se concluye que la integración y la evaluación agrupada son reproducibles, mientras que la predicción disponible es insuficiente para priorizar personas. Se recomienda verificar localmente las barreras identificadas, recolectar trayectoria temporal y nacimiento con referencia escolar, y validar en otra cohorte antes de evaluar intervenciones. Se conserva el resultado negativo de predicción como evidencia, sin añadir edad o grado para inflar métricas.')
h('Referencias')
refs=[
'[1] INEI. ENAHO panel 2021–2025. Diccionario de datos, secciones del corte 2025, pp. 588–594 y 657–666. ZIP originales 1041, módulos 1474, 1479, 1475 y 1478. Procedencia y hashes en docs/manifiesto_fuentes.json.',
'[2] MINEDU ESCALE. Padron.dbf y Padlocal.dbf. Padrón de servicios educativos, actualización 29 de octubre de 2025. Distribución original suministrada para el curso.',
'[3] MINEDU. Edades normativas. Referencia al 31 de marzo. https://www.gob.pe/institucion/minedu/informes-publicaciones/2742610-edades-normativas',
'[4] MINEDU ESCALE. Alumnos con atraso escolar. Metadatos del indicador. https://escale.minedu.gob.pe/tendencias-2016-portlet/servlet/tendencias/archivo?idCuadro=14&tipo=meta',
'[5] scikit-learn. Common pitfalls and recommended practices. https://scikit-learn.org/stable/common_pitfalls',
'[6] S. Lundberg y S. I. Lee. A Unified Approach to Interpreting Model Predictions. NeurIPS 30, 2017. https://proceedings.neurips.cc/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html']
for ref in refs:
 x=p(ref)
 for rr in x.runs:rr.font.size=Pt(8)
 x.paragraph_format.space_after=Pt(3)
path=R/'informe/Informe_Jose_Guillermo_Ticona_Lipa.docx';doc.save(path)
(R/'informe/informe_fuente.md').write_text('\n\n'.join(source),encoding='utf8');print(path)
